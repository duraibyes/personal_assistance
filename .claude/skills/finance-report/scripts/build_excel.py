"""Build an income & expense Excel workbook from the JSON written by export-data.mjs.

Usage: python build_excel.py <data.json> [out.xlsx]
Requires openpyxl (pip install openpyxl).
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = '"₹"#,##0.00;[Red]-"₹"#,##0.00'
DATE = "dd-mmm-yyyy"
PERCENT = "0.0%"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="thin"))


def write_table(ws, top, left, headers, rows, money=(), dates=(), percents=(), total_label="Total"):
    """Write a header, data rows and an optional SUM total row. Returns (first_row, last_row) of the data."""
    for j, header in enumerate(headers):
        cell = ws.cell(row=top, column=left + j, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = ws.cell(row=top + i, column=left + j, value=value)
            if j in money:
                cell.number_format = MONEY
            elif j in dates:
                cell.number_format = DATE
            elif j in percents:
                cell.number_format = PERCENT
    first, last = top + 1, top + len(rows)
    if total_label:
        total_row = last + 1
        ws.cell(row=total_row, column=left, value=total_label).font = BOLD
        for j in money:
            col = get_column_letter(left + j)
            cell = ws.cell(row=total_row, column=left + j, value=f"=SUM({col}{first}:{col}{last})" if rows else 0)
            cell.number_format, cell.font = MONEY, BOLD
        for j in range(len(headers)):
            ws.cell(row=total_row, column=left + j).border = TOTAL_BORDER
    return first, last


def autosize(ws, min_width=8, max_width=48):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            numeric = isinstance(cell.value, (int, float, date)) or str(cell.value).startswith("=")
            length = 15 if numeric else len(str(cell.value))
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[col].width = max(min_width, min(max_width, width + 2))


def detail_sheet(wb, title, headers, rows, money, dates):
    ws = wb.create_sheet(title)
    write_table(ws, 1, 1, headers, rows, money=money, dates=dates)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    autosize(ws)


def summary_sheet(ws, data):
    meta, s = data["meta"], data["summary"]
    ws.title = "Summary"
    ws["A1"] = "Income & Expense Report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{meta['from']} to {meta['to']}  ·  {meta['scope']}"
    generated = datetime.fromisoformat(meta["generatedAt"].replace("Z", "+00:00")).astimezone()
    ws["A3"] = f"Generated {generated:%d %b %Y %H:%M}"
    ws["A3"].font = Font(italic=True, color="6B7280")

    kpis = [
        ("Total income", s["totalIncome"]),
        ("Expenses", s["totalExpenses"]),
        ("Loan EMIs", s["totalEmi"]),
        ("Total outflow", s["totalOutflow"]),
        ("Net savings", s["net"]),
    ]
    for i, (label, value) in enumerate(kpis, start=5):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=value).number_format = MONEY
    ws.cell(row=10, column=1, value="Savings rate").font = BOLD
    rate = ws.cell(row=10, column=2, value=s["savingsRate"] / 100 if s["savingsRate"] is not None else "n/a")
    rate.number_format = PERCENT

    top = 12
    rows = [[m["label"], m["income"], m["expenses"], m["emi"], None] for m in data["monthly"]]
    first, last = write_table(ws, top, 1, ["Month", "Income", "Expenses", "Loan EMIs", "Net"], rows, money=(1, 2, 3, 4))
    for r in range(first, last + 1):
        ws.cell(row=r, column=5, value=f"=B{r}-C{r}-D{r}").number_format = MONEY

    if rows:
        chart = BarChart()
        chart.type, chart.grouping = "col", "clustered"
        chart.title, chart.height, chart.width = "Income vs outflow by month", 8, 18
        chart.y_axis.numFmt = '"₹"#,##0'
        chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=top, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        ws.add_chart(chart, "G4")
    autosize(ws)
    ws.column_dimensions["A"].width = 18


def breakdown_sheet(wb, data):
    ws = wb.create_sheet("Breakdown")
    sections = [
        ("Expenses by category", "Category", data["expenseByCategory"]),
        ("Income by source", "Source", data["incomeBySource"]),
        ("Expenses by payment method", "Payment method", data["expenseByPaymentMethod"]),
        ("Loan EMIs by loan", "Loan", data["emiByLoan"]),
    ]
    top = 1
    for title, label, groups in sections:
        ws.cell(row=top, column=1, value=title).font = Font(bold=True, size=13)
        rows = [[g["name"], g["total"], g["count"], g["share"] / 100] for g in groups]
        first, last = write_table(ws, top + 1, 1, [label, "Amount", "Entries", "Share"], rows, money=(1,), percents=(3,))
        if title == "Expenses by category" and rows:
            pie = PieChart()
            pie.title, pie.height, pie.width = title, 8, 12
            pie.add_data(Reference(ws, min_col=2, min_row=first, max_row=last))
            pie.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
            ws.add_chart(pie, f"G{top}")
        top = max(last + 4, top + 18 if title == "Expenses by category" and rows else 0)
    autosize(ws)


def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python build_excel.py <data.json> [out.xlsx]")
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".xlsx")
    data = json.loads(src.read_text(encoding="utf-8"))

    wb = Workbook()
    summary_sheet(wb.active, data)
    breakdown_sheet(wb, data)
    detail_sheet(
        wb,
        "Income",
        ["Date", "Amount", "Source", "Category", "Description"],
        [[date.fromisoformat(r["date"]), r["amount"], r["source"], r["category"], r["description"]] for r in data["incomes"]],
        money=(1,),
        dates=(0,),
    )
    detail_sheet(
        wb,
        "Expenses",
        ["Date", "Amount", "Category", "Description", "Payment method", "Vendor"],
        [
            [date.fromisoformat(r["date"]), r["amount"], r["category"], r["description"], r["paymentMethod"], r["vendor"]]
            for r in data["expenses"]
        ],
        money=(1,),
        dates=(0,),
    )
    if data["meta"].get("includesEmi", True):
        detail_sheet(
            wb,
            "Loan EMIs",
            ["Date", "Amount", "Loan", "Lender", "EMI #", "Type", "Note"],
            [
                [date.fromisoformat(r["date"]), r["amount"], r["loan"], r["lender"], r["emiNumber"], r["type"], r["note"]]
                for r in data["emis"]
            ],
            money=(1,),
            dates=(0,),
        )

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(out.resolve())


if __name__ == "__main__":
    main()
